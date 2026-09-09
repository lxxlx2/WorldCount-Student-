from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")
        tcMar.append(node)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def style_document(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.4), Cm(2.2)
    sec.left_margin, sec.right_margin = Cm(2.6), Cm(2.4)
    sec.header_distance, sec.footer_distance = Cm(1.2), Cm(1.2)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Noto Sans SC"; normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans SC")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    for name, size, font in (("Title", 22, "Noto Sans SC"), ("Heading 1", 16, "Noto Sans SC"), ("Heading 2", 15, "Noto Sans SC"), ("Heading 3", 14, "Noto Sans SC")):
        s = styles[name]; s.font.name = font; s.font.size = Pt(size); s.font.color.rgb = RGBColor(0,0,0)
        s._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        s.paragraph_format.space_before = Pt(12); s.paragraph_format.space_after = Pt(8)
    add_page_number(sec.footer.paragraphs[0])


def title_page(doc, course_code, course_name, subject):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(55)
    r = p.add_run("四川大学高等教育自学考试"); r.bold = True; r.font.size = Pt(22); r.font.name = "Noto Sans SC"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans SC")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(35)
    r = p.add_run("实践报告"); r.bold = True; r.font.size = Pt(30); r.font.name = "Noto Sans SC"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans SC")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(28)
    r = p.add_run(subject); r.bold = True; r.font.size = Pt(18); r.font.name = "Noto Sans SC"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans SC")
    doc.add_paragraph("")
    table = doc.add_table(rows=6, cols=2); table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rows = [("专业", "【请填写专业名称及专业代码】"), ("层次", "本科"), ("课程代码", course_code),
            ("课程名称", course_name), ("准考证号", "【请填写准考证号】"), ("考生姓名", "【请填写姓名】")]
    for row, (a,b) in zip(table.rows, rows):
        row.cells[0].text=a; row.cells[1].text=b
        row.cells[0].width=Cm(4); row.cells[1].width=Cm(9)
        for c in row.cells:
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(c,150,140,150,140)
            for pp in c.paragraphs: pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(48)
    p.add_run("学历教育办公室制\n二〇二六年九月").font.size=Pt(14)
    doc.add_page_break()


def integrity(doc):
    doc.add_heading("考生诚信承诺书", 0)
    p=doc.add_paragraph()
    p.paragraph_format.first_line_indent=Cm(0.85)
    p.add_run("本人郑重承诺：本次课程实践考核报告，由本人在理解考核要求的基础上独立完成。报告中所涉及的实践过程、数据和分析结论均应由本人复核并确认真实有效。本人提交前将逐项检查代码、日志、截图和文字内容，并能够在抽查答辩中解释其实现原理与操作过程。")
    p=doc.add_paragraph("如有违反以上承诺，愿承担包括但不限于考核成绩作废、纪律处分等一切责任。")
    p.paragraph_format.first_line_indent=Cm(0.85)
    doc.add_paragraph("\n\n承诺人：________________\n\n日期：______年____月____日")
    doc.add_paragraph("提交前必须由考生本人核对内容并亲笔签署。")
    doc.add_page_break()


def add_toc(doc, items):
    doc.add_heading("目录", 0)
    for i, item in enumerate(items,1):
        p=doc.add_paragraph(f"{i}  {item}")
        p.paragraph_format.space_after=Pt(10)
    doc.add_page_break()


def add_table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style="Table Grid"
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=str(h); set_cell_shading(c,"1F4E79")
        for r in c.paragraphs[0].runs: r.font.color.rgb=RGBColor(255,255,255); r.bold=True
        c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
    for ridx,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            if ridx%2: set_cell_shading(cells[i],"F2F6FA")
            cells[i].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 or len(str(val))<18 else WD_ALIGN_PARAGRAPH.LEFT
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Cm(w)
    doc.add_paragraph("")
    return t


def add_code(doc, code, max_lines=None):
    if max_lines: code="\n".join(code.splitlines()[:max_lines])
    p=doc.add_paragraph()
    p.paragraph_format.left_indent=Cm(0.5); p.paragraph_format.right_indent=Cm(0.5)
    p.paragraph_format.line_spacing=1.0; p.paragraph_format.space_after=Pt(8)
    r=p.add_run(code); r.font.name="Menlo"; r.font.size=Pt(8.5)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans SC")


def add_figure(doc, path, caption, width=15.5):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(width))
    p=doc.add_paragraph(caption); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=False


def database_report():
    doc=Document(); style_document(doc)
    title_page(doc,"08695","数据库及其应用操作（实践）","题目一 学生选课管理系统")
    integrity(doc)
    add_toc(doc,["实践题目与开发环境","需求分析","数据库设计","SQL 操作实现","运行截图与结果分析","实践问题及解决方法","实践总结","参考资料","提交前检查"])
    doc.add_heading("1 实践题目与开发环境",1)
    doc.add_paragraph("所选题目：题目一 学生选课管理系统。系统围绕学生、课程及选课成绩建立关系数据库，支持基础数据维护、成绩自动计算、关联查询、分组统计、视图访问和索引优化。")
    add_table(doc,["项目","实际环境"],[("操作系统","macOS 26.6.2"),("数据库管理系统","MySQL Community Server 8.4.11"),("数据库管理工具","MySQL 命令行客户端"),("数据库名称","student_course_management"),("字符集","utf8mb4")],[4,11])
    doc.add_heading("2 需求分析",1)
    doc.add_paragraph("本系统解决学生基本信息、课程目录和选课成绩分散管理的问题。教务人员负责维护学生和课程，任课教师录入平时与期末成绩，管理人员通过查询和统计了解课程成绩与院系选课情况。主要数据包括学生身份与院系、课程学分与类别、学生和课程之间的选课关联及成绩。")
    doc.add_paragraph("功能包括：维护三类基础数据；用主键、外键、唯一约束和检查约束保证完整性；以生成列统一计算总评成绩；完成单表、多表连接、聚合和条件查询；通过视图提供常用结果；在成绩表学号字段上建立索引。")
    doc.add_heading("3 数据库设计",1)
    doc.add_heading("3.1 E-R 图",2)
    add_figure(doc,ROOT/"database_student_course/evidence/fig03_er_diagram.png","图 1 学生选课管理系统 E-R 图",15.5)
    doc.add_heading("3.2 关系模式",2)
    doc.add_paragraph("学生（学号【PK】，姓名，性别，出生日期，所在院系，入学年份，联系电话）")
    doc.add_paragraph("课程（课程号【PK】，课程名，学分，学时，课程类别，开课院系）")
    doc.add_paragraph("选课成绩（选课编号【PK】，学号【FK】，课程号【FK】，平时成绩，期末成绩，总评成绩，选课日期）")
    doc.add_paragraph("学生和课程是多对多关系，由选课成绩关系进行分解；tblStudent 到 tblScore、tblCourse 到 tblScore 均为 1:N。三张表字段均只保存原子值，非主属性完全依赖候选键，不存在非主属性之间的传递依赖，满足第三范式。")
    doc.add_heading("3.3 数据表结构",2)
    add_table(doc,["表","字段","类型","约束与说明"],[
        ("tblStudent","student_id","CHAR(10)","主键，学号"),("","student_name","VARCHAR(30)","非空，姓名"),("","gender","ENUM","男或女"),("","birth_date","DATE","出生日期"),("","department","VARCHAR(50)","所在院系"),("","enrollment_year","YEAR","入学年份"),("","phone","VARCHAR(20)","唯一，联系电话"),
        ("tblCourse","course_id","CHAR(6)","主键，课程号"),("","course_name","VARCHAR(60)","非空，课程名"),("","credits","TINYINT","1 至 6 学分"),("","class_hours","SMALLINT","学时大于 0"),("","course_type","ENUM","必修、选修或公选"),("","offering_department","VARCHAR(50)","开课院系"),
        ("tblScore","enrollment_id","INT", "自增主键"),("","student_id","CHAR(10)","外键，级联更新"),("","course_id","CHAR(6)","外键，级联更新"),("","usual_score","TINYINT","0 至 100"),("","final_score","TINYINT","0 至 100"),("","total_score","DECIMAL(5,2)","生成列，30%+70%"),("","selected_at","DATE","选课日期")],[2.8,3.8,3.3,6.2])
    doc.add_heading("4 SQL 操作实现",1)
    doc.add_heading("4.1 数据定义 DDL",2)
    doc.add_paragraph("脚本 01_schema.sql 创建数据库、三张 InnoDB 表、完整性约束、索引和视图。以下为核心定义，完整可执行脚本随程序文件提交。")
    add_code(doc,(ROOT/"database_student_course/sql/01_schema.sql").read_text(),max_lines=90)
    doc.add_heading("4.2 数据操纵 DML",2)
    doc.add_paragraph("02_seed.sql 插入 10 名学生、10 门课程和 21 条有效选课记录。03_operations_and_queries.sql 将 Web 应用开发课程调整为 4 学分，并通过插入后删除一条专用演示记录验证 DELETE；这样既覆盖 DML，又不破坏最终业务数据。")
    add_code(doc,"UPDATE tblCourse SET credits = 4 WHERE course_id = 'CS0005';\nINSERT INTO tblScore(student_id, course_id, usual_score, final_score, selected_at)\nVALUES ('2026000010','GE0001',80,82,'2026-09-03');\nDELETE FROM tblScore WHERE student_id='2026000010' AND course_id='GE0001';")
    doc.add_heading("4.3 数据查询 DQL",2)
    doc.add_paragraph("项目实现 6 类查询：学生列表、数据库原理课程成绩、每生选课门数与平均成绩、90 分及以上名单、院系选课人数、视图查询。查询同时覆盖 ORDER BY、INNER/LEFT JOIN、COUNT、AVG、GROUP BY、DISTINCT 和条件过滤。")
    add_code(doc,(ROOT/"database_student_course/sql/03_operations_and_queries.sql").read_text(),max_lines=100)
    doc.add_heading("4.4 视图与索引",2)
    add_code(doc,"CREATE INDEX idx_score_student_id ON tblScore(student_id);\n\nCREATE VIEW vw_student_course_score AS\nSELECT s.student_id, s.student_name, c.course_name, sc.total_score\nFROM tblScore sc\nJOIN tblStudent s ON s.student_id=sc.student_id\nJOIN tblCourse c ON c.course_id=sc.course_id;")
    doc.add_heading("5 运行截图与结果分析",1)
    ev=ROOT/"database_student_course/evidence"
    figures=[("fig01_tables_structure.png","图 2 数据库所有表及字段结构"),("fig02_table_data.png","图 3 三张表测试数据"),("fig03_er_diagram.png","图 4 表间关系示意"),("fig04_database_course.png","图 5 数据库原理课程成绩查询"),("fig05_student_summary.png","图 6 每名学生选课门数与平均成绩"),("fig06_high_scores.png","图 7 总评成绩 90 分及以上名单"),("fig07_department_count.png","图 8 各院系选课人数统计"),("fig08_view_index.png","图 9 视图结果与索引验证")]
    for f,c in figures:
        add_figure(doc,ev/f,c)
        if f=="fig02_table_data.png": doc.add_paragraph("最终数据量为学生 10 行、课程 10 行、选课成绩 21 行，满足每张表不少于 10 条测试数据的要求。")
    doc.add_paragraph("04_verification.sql 的执行结果显示：成绩范围违规数 0、外键孤儿记录数 0、视图行数 21；idx_score_student_id 确认位于 tblScore.student_id。")
    doc.add_heading("6 实践问题及解决方法",1)
    doc.add_heading("6.1 总评成绩一致性",2)
    doc.add_paragraph("最初若在插入数据时人工填写总评成绩，平时或期末成绩更新后容易出现不一致。解决方法是使用 MySQL STORED 生成列，根据 ROUND(usual_score*0.30+final_score*0.70,2) 自动计算，避免重复维护。")
    doc.add_heading("6.2 重复选课与外键完整性",2)
    doc.add_paragraph("仅设置自增主键仍可能让同一学生重复选择同一课程，因此增加 (student_id, course_id) 唯一约束；两个外键使用 ON UPDATE CASCADE 和 ON DELETE RESTRICT，既支持编号更新，也防止误删仍被成绩表引用的主数据。")
    doc.add_heading("6.3 删除操作影响正式数据",2)
    doc.add_paragraph("直接删除已有成绩会改变查询结果。实践中先插入一条专门用于演示的记录，再按学生和课程精确删除，使 DELETE 得到真实执行，同时最终数据集保持可重复验证。")
    doc.add_heading("7 实践总结",1)
    doc.add_paragraph("本次实践从需求分析开始，把学生、课程与选课成绩分解为三个满足第三范式的关系，随后用主键、外键、唯一约束和检查约束落实数据规则。我进一步掌握了生成列在派生数据一致性方面的作用，并通过单表查询、多表连接、聚合统计、视图和索引把逻辑设计转化为可执行数据库。实际运行时，先重建数据库，再装载数据、执行维护操作和查询，最后用独立验证脚本检查行数、成绩范围、孤儿记录、视图和索引，使结果可以重复。当前不足是项目以命令行脚本为主，尚未开发图形化业务界面；若继续扩展，可增加用户权限、成绩变更审计、事务处理、备份恢复和 Web 管理端。")
    doc.add_heading("8 参考资料",1)
    doc.add_paragraph("[1] 王珊、萨师煊：《数据库系统概论》第 6 版，高等教育出版社，2023 年。\n[2] MySQL 8.4 Reference Manual，CREATE TABLE、Generated Column、CREATE VIEW 与 CREATE INDEX 章节。\n[3] 08695《数据库及其应用操作（实践）》实践考核方案及实践报告模板，2026 年。")
    doc.add_heading("9 提交前检查",1)
    add_table(doc,["材料","状态或操作"],[("实践报告 PDF","已生成；提交前按专业代码_08695_准考证号_姓名.pdf 重命名"),("SQL 脚本","已完成并实际执行"),("README.txt","已完成"),("身份信息与承诺书","考生本人填写并签署"),("知网查重及 AIGC 检测","需考生本人完成并按通知打包")],[5,10])
    doc._body._body.remove(doc.paragraphs[-1]._element)
    path=OUT/"08695_题目一_学生选课管理系统_实践报告_待填身份.docx"; doc.save(path); return path


def bigdata_report():
    doc=Document(); style_document(doc)
    title_page(doc,"14899","大数据技术基础（实践）","题目三 Hadoop 平台搭建与 WordCount")
    integrity(doc)
    add_toc(doc,["实践题目与开发环境","环境准备","Hadoop 安装与配置","Hadoop 启动与验证","HDFS 操作","MapReduce WordCount 运行","运行结果汇总","实践问题及解决方法","实践总结","参考资料","提交前检查"])
    doc.add_heading("1 实践题目与开发环境",1)
    doc.add_paragraph("所选题目：题目三 Hadoop 平台搭建与 MapReduce 示例运行。实践完成 HDFS 单节点服务配置、NameNode 和 DataNode 启动、HDFS 文件操作，以及 Hadoop 自带 WordCount 示例的实际执行和结果统计。")
    add_table(doc,["项目","实际环境"],[("主机操作系统","macOS 26.6.2，Unix 命令行环境"),("Hadoop 版本","Apache Hadoop 3.5.0"),("JDK 版本","OpenJDK 17.0.20.1"),("HDFS 模式","单节点伪分布式，副本数 1"),("MapReduce 模式","local 本地执行模式"),("HDFS 地址","hdfs://localhost:9000"),("NameNode Web","http://localhost:9870")],[4,11])
    doc.add_paragraph("说明：本次实操在 macOS 的 Unix 环境真实完成，而考核模板建议 Ubuntu 或 CentOS。Hadoop、HDFS 和 MapReduce 命令及配置文件保持一致；如考点严格限定 Linux，可按 README 在 Linux 复现并替换图 1。")
    doc.add_heading("2 环境准备",1)
    doc.add_paragraph("通过 Homebrew 安装 Hadoop 3.5.0，其依赖提供 OpenJDK 17。运行脚本显式设置 JAVA_HOME、HADOOP_HOME、HADOOP_CONF_DIR 和 PATH，避免系统其他 JDK 版本影响 Hadoop。")
    add_figure(doc,ROOT/"bigdata_wordcount/evidence/01_linux_java.png","图 1 操作系统与 JDK 版本验证")
    add_figure(doc,ROOT/"bigdata_wordcount/evidence/02_hadoop_version.png","图 2 Hadoop 安装目录与版本验证")
    doc.add_heading("3 Hadoop 安装与配置",1)
    doc.add_paragraph("core-site.xml 将默认文件系统指向本机 NameNode 的 9000 端口；hdfs-site.xml 将副本数设为 1，并把 NameNode、DataNode 数据目录限定到本实践专用目录；mapred-site.xml 使用 local 模式运行 MapReduce，便于在单机环境观察 map、shuffle 和 reduce 流程。")
    add_figure(doc,ROOT/"bigdata_wordcount/evidence/03_configuration.png","图 3 Hadoop 核心配置文件")
    doc.add_heading("4 Hadoop 启动与验证",1)
    doc.add_paragraph("首次运行通过 hdfs namenode -format -force -nonInteractive 格式化专用 NameNode；随后分别启动 NameNode 和 DataNode。脚本等待安全模式退出后再执行写操作，避免刚启动时删除或上传失败。jps 显示两个 HDFS Java 进程，dfsadmin 报告显示 1 个存活 DataNode。")
    add_figure(doc,ROOT/"bigdata_wordcount/evidence/04_processes.png","图 4 Hadoop 进程与 HDFS 容量验证")
    add_figure(doc,ROOT/"bigdata_wordcount/evidence/10_hdfs_web.png","图 5 NameNode Web 管理界面")
    doc.add_heading("5 HDFS 操作",1)
    doc.add_paragraph("在 /user/student_demo 下创建 input 目录，上传三个输入文件，并用 ls 与 cat 验证文件数量和内容。运行脚本还会创建 output、下载 part-r-00000，并在重复执行前删除本项目自己的旧 input 和 output，从而保持幂等。正式提交前可通过 STUDENT_ID=准考证号 将个人目录替换为要求的 /user/准考证号。")
    add_figure(doc,ROOT/"bigdata_wordcount/evidence/05_hdfs_input.png","图 6 创建 HDFS 目录并上传三个文件")
    add_figure(doc,ROOT/"bigdata_wordcount/evidence/06_input_content.png","图 7 查看 HDFS 文件内容和本地行数")
    doc.add_heading("6 MapReduce WordCount 运行",1)
    doc.add_paragraph("项目自动查找 hadoop-mapreduce-examples-3.5.0.jar，并以 wordcount 作为示例类执行。三个输入文件对应三个 map split；Mapper 共输出 73 条键值记录，Combiner 在本地预聚合，Reducer 接收 46 个不同键并输出 46 行。日志中 map 和 reduce 均达到 100%，作业状态为 completed successfully，Failed Shuffles 为 0。")
    add_figure(doc,ROOT/"bigdata_wordcount/evidence/07_wordcount_job.png","图 8 WordCount 作业命令、计数器与成功状态")
    add_figure(doc,ROOT/"bigdata_wordcount/evidence/08_result.png","图 9 HDFS 输出文件 part-r-00000 内容")
    add_figure(doc,ROOT/"bigdata_wordcount/evidence/09_summary_top5.png","图 10 单词总数、不同单词数与 Top 5")
    doc.add_heading("7 运行结果汇总",1)
    add_table(doc,["序号","操作","命令或对象","实际结果"],[(1,"验证版本","hadoop version","Hadoop 3.5.0"),(2,"启动 HDFS","hdfs --daemon start","NameNode、DataNode 存活"),(3,"创建目录","hdfs dfs -mkdir -p","/user/student_demo/input"),(4,"上传文件","hdfs dfs -put","3 个文件，共 12 行"),(5,"运行 WordCount","hadoop jar ... wordcount","map 100%，reduce 100%"),(6,"查看结果","hdfs dfs -cat","46 个不同单词"),(7,"汇总词数","awk","总词数 73"),(8,"Top 5","sort -k2,2nr","data 8，mapreduce 5，hadoop 4，hdfs 3，in 3"),(9,"下载结果","hdfs dfs -get -f","downloaded-part-r-00000 已生成")],[1.3,3.2,4.5,6.2])
    doc.add_heading("8 实践问题及解决方法",1)
    doc.add_heading("8.1 JDK 版本继承错误",2)
    doc.add_paragraph("系统原有 JAVA_HOME 指向 JDK 25，首次环境日志因此显示了非预期版本。解决方法是不再直接继承 JAVA_HOME，而让运行脚本默认固定到 Hadoop 3.5.0 的 Homebrew 依赖 JDK 17，并保留 HADOOP_JAVA_HOME 作为显式覆盖入口。重新运行后日志确认 OpenJDK 17.0.20.1。")
    doc.add_heading("8.2 NameNode 安全模式导致清理失败",2)
    doc.add_paragraph("NameNode 重启后处于安全模式扩展期，此时 rm 无法删除旧 input/output，后续 put 报 File exists。脚本增加 hdfs dfsadmin -safemode wait，确认安全模式退出后再清理旧目录，使同一流程可以重复执行。")
    doc.add_heading("8.3 命令警告污染结果文件",2)
    doc.add_paragraph("macOS 没有 Hadoop 本机原生库，hdfs dfs -cat 会输出 NativeCodeLoader 警告，直接通过管道保存可能污染 part-r-00000。解决方法是先用 hdfs dfs -get 下载 HDFS 原始文件，再从下载文件复制和统计。最终 part-r-00000 仅包含 46 行“单词-次数”数据。")
    doc.add_heading("9 实践总结",1)
    doc.add_paragraph("本次实践完成了从 Hadoop 环境配置、HDFS 服务启动、输入数据上传到 MapReduce WordCount 处理和结果分析的完整流程。我通过 core-site.xml 理解客户端如何定位 NameNode，通过 hdfs-site.xml 设置单节点副本策略，并从 jps、dfsadmin 和 Web 页面三个角度确认集群状态。WordCount 的计数器把处理过程量化为 3 个输入分片、12 条输入记录、73 条 Map 输出和 46 条 Reduce 输出，使 Mapper、Combiner、Shuffle 与 Reducer 的职责更直观。实际排错还说明环境变量、安全模式和标准输出处理会直接影响可重复性。当前实践仅使用一个 DataNode，不能体现跨节点容错和网络传输；后续可在多台 Linux 虚拟机上部署全分布式集群，并增加 YARN、资源调度、数据倾斜和故障恢复实验。")
    doc.add_heading("10 参考资料",1)
    doc.add_paragraph("[1] 林子雨：《大数据基础编程、实验和案例教程》第 3 版，清华大学出版社，2024 年。\n[2] Apache Hadoop 3.5.0 Documentation，HDFS Commands Guide 与 MapReduce Tutorial。\n[3] 14899《大数据技术基础（实践）》实践考核方案及实践报告模板，2026 年。")
    doc.add_heading("11 提交前检查",1)
    add_table(doc,["材料","状态或操作"],[("实践报告 PDF/Word","已生成；提交前按专业代码_14899_准考证号_姓名重命名"),("配置文件、输入、输出和日志","已生成并纳入程序包"),("README.txt","已完成"),("个人 HDFS 目录","提交前用 STUDENT_ID=准考证号 重新运行"),("Linux 截图","若考点严格限定 Linux，应在 Linux 复现后替换图 1"),("身份信息与承诺书","考生本人填写并签署"),("知网查重及 AIGC 检测","需考生本人完成并按通知打包")],[5,10])
    doc._body._body.remove(doc.paragraphs[-1]._element)
    path=OUT/"14899_题目三_WordCount_实践报告_待填身份.docx"; doc.save(path); return path


if __name__ == "__main__":
    print(database_report())
    print(bigdata_report())
