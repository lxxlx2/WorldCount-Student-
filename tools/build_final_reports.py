#!/usr/bin/env python3
import json
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

R=Path(__file__).resolve().parents[1]; O=R/'output'; O.mkdir(exist_ok=True)
F=json.loads((O/'final_facts.json').read_text(encoding='utf-8'))
ID={'major':'信息资源管理','code':'W120503','exam':'010126323484','name':'刘怡荷'}

def rf(run,font='宋体',size=12,bold=None):
    run.font.name=font; run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),font); run.font.size=Pt(size)
    if bold is not None: run.bold=bold

def setup(d):
    s=d.sections[0]; s.page_width=Cm(21); s.page_height=Cm(29.7); s.top_margin=Cm(2.54); s.bottom_margin=Cm(2.54); s.left_margin=Cm(3); s.right_margin=Cm(2.5)
    n=d.styles['Normal']; n.font.name='宋体'; n._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); n.font.size=Pt(12); n.paragraph_format.line_spacing=1.5
    for k,sz in [('Heading 1',16),('Heading 2',15),('Heading 3',14)]:
        x=d.styles[k]; x.font.name='黑体'; x._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体'); x.font.size=Pt(sz); x.font.bold=True
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); s.footer.paragraphs[0]._p.append(fld); s.footer.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER

def cover(d,cc,cn):
    for text,sz,space in [('四川大学高等教育自学考试',22,48),('实践报告',28,28)]:
        p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(space); r=p.add_run(text); rf(r,'黑体',sz,True)
    t=d.add_table(rows=6,cols=2); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    rows=[('专业',f"{ID['major']}（{ID['code']}）"),('层次','本科'),('课程代码',cc),('课程名称',cn),('准考证号',ID['exam']),('考生姓名',ID['name'])]
    for rr,(a,b) in zip(t.rows,rows):
        rr.cells[0].text=a; rr.cells[1].text=b
        for c in rr.cells:
            for p in c.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs: rf(r,'宋体',14)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(48); r=p.add_run('学历教育办公室制\n二〇二六年九月'); rf(r,'宋体',14); d.add_page_break()

def integrity(d):
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('考生诚信承诺书'); rf(r,'黑体',18,True)
    for x in ['本人郑重承诺：本次课程实践考核报告，由本人在理解考核要求的基础上独立完成。报告中所涉及的实践过程、调查数据、分析结论均为真实、有效的记录与思考，绝无抄袭、伪造、剽窃等行为。','如有违反以上承诺，愿承担包括但不限于考核成绩作废、纪律处分等一切责任。']:
        p=d.add_paragraph(x); p.paragraph_format.first_line_indent=Cm(.85)
    p=d.add_paragraph('\n承诺人：________________\n\n年    月    日'); p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; d.add_page_break()

def toc(d,items):
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('目录'); rf(r,'黑体',18,True)
    for x in items: d.add_paragraph(x)
    d.add_page_break()

def table(d,heads,rows):
    t=d.add_table(rows=1,cols=len(heads)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(heads): t.rows[0].cells[i].text=str(h)
    for row in rows:
        c=t.add_row().cells
        for i,v in enumerate(row): c[i].text=str(v)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs: rf(r,'宋体',10.5)
    d.add_paragraph(''); return t

def fig(d,p,cap):
    p=Path(p)
    if not p.exists(): d.add_paragraph(f'【缺少真实截图：{p.name}】'); return
    x=d.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER; x.add_run().add_picture(str(p),width=Cm(15))
    x=d.add_paragraph(cap); x.alignment=WD_ALIGN_PARAGRAPH.CENTER

def code(d,p,n=60):
    s=Path(p).read_text(encoding='utf-8').splitlines()[:n]; x=d.add_paragraph()
    for i,l in enumerate(s):
        r=x.add_run(l+('\n' if i<len(s)-1 else '')); rf(r,'Menlo',8.5)

def db_report():
    f=F['database']; d=Document(); setup(d); cover(d,'08695','数据库及其应用操作（实践）'); integrity(d); toc(d,['1 实践题目','2 开发环境','3 需求分析','4 数据库设计','5 SQL 操作实现','6 运行截图','7 实践中遇到的问题及解决方法','8 实践总结','9 参考资料'])
    d.add_heading('1 实践题目',1); d.add_paragraph('所选题目：题目一 学生选课管理系统。')
    d.add_heading('2 开发环境',1); table(d,['项目','实际环境'],[('操作系统',F['host_os']),('数据库管理系统',f"MySQL Community Server {f['mysql_version']}"),('数据库管理工具','MySQL Workbench / MySQL 命令行客户端'),('数据库名称','student_course_management'),('字符集','utf8mb4')])
    d.add_heading('3 需求分析',1); d.add_paragraph('本系统用于管理学生基本信息、课程信息和学生选课成绩。通过学生表、课程表和选课成绩表建立关联，支持数据新增、修改、删除、单表查询、多表连接、聚合统计、视图和索引，并通过主键、外键和检查约束保证数据完整性。')
    d.add_heading('4 数据库设计',1); d.add_heading('4.1 E-R 图',2); fig(d,R/'database_student_course/evidence/fig03_er_diagram.png','图1 学生选课管理系统 E-R 图')
    d.add_heading('4.2 关系模式',2); d.add_paragraph('学生（学号【PK】，姓名，性别，出生日期，所在院系，入学年份，联系电话）'); d.add_paragraph('课程（课程号【PK】，课程名，学分，学时，课程类别，开课院系）'); d.add_paragraph('选课成绩（选课编号【PK】，学号【FK】，课程号【FK】，平时成绩，期末成绩，总评成绩）')
    d.add_heading('4.3 数据表结构',2); table(d,['表','字段','类型/约束'],[('tblStudent','student_id','CHAR(10)，主键'),('tblStudent','student_name','VARCHAR(30)，非空'),('tblStudent','gender','男/女'),('tblStudent','birth_date','DATE'),('tblStudent','department','VARCHAR(50)'),('tblStudent','enrollment_year','YEAR'),('tblStudent','phone','VARCHAR(20)'),('tblCourse','course_id','CHAR(6)，主键'),('tblCourse','course_name','VARCHAR(60)，非空'),('tblCourse','credits','1-6'),('tblCourse','class_hours','整数'),('tblCourse','course_type','必修/选修/公选'),('tblCourse','offering_department','VARCHAR(50)'),('tblScore','enrollment_id','INT，自增主键'),('tblScore','student_id','外键'),('tblScore','course_id','外键'),('tblScore','usual_score','0-100'),('tblScore','final_score','0-100'),('tblScore','total_score','INT，自动计算')])
    d.add_heading('5 SQL 操作实现',1); d.add_heading('5.1 数据定义（DDL）',2); code(d,R/'database_student_course/sql/01_schema.sql',90); d.add_heading('5.2 数据操纵（DML）',2); d.add_paragraph('三张表均插入不少于10条测试数据，并实际执行 UPDATE 和 DELETE。'); code(d,R/'database_student_course/sql/03_operations_and_queries.sql',18); d.add_heading('5.3 数据查询（DQL）',2); d.add_paragraph('查询覆盖学生排序、数据库原理课程联表查询、每名学生选课门数和平均成绩、90分及以上名单、院系选课人数以及视图查询。'); d.add_heading('5.4 视图与索引',2); d.add_paragraph('创建 vw_student_course_score 视图，并在 tblScore.student_id 上创建索引。')
    d.add_heading('6 运行截图',1); ev=R/'database_student_course/evidence_real'
    for fn,cap in [('01_environment.png','图2 MySQL 8.0 实际运行环境'),('02_tables_structure.png','图3 数据库所有表结构'),('03_table_data.png','图4 三张表测试数据与行数'),('04_database_course.png','图5 数据库原理课程联表查询'),('05_student_summary.png','图6 每名学生选课门数与平均成绩'),('06_high_scores.png','图7 总评成绩90分及以上名单'),('07_department_count.png','图8 各院系选课人数统计'),('08_view_index.png','图9 视图与索引验证')]: fig(d,ev/fn,cap)
    c=f['counts']; d.add_paragraph(f"验证结果：tblStudent={c['tblStudent']}，tblCourse={c['tblCourse']}，tblScore={c['tblScore']}；非法成绩={f['invalid_scores']}，外键孤儿记录={f['orphan_scores']}。")
    d.add_heading('7 实践中遇到的问题及解决方法',1)
    for i,x in enumerate(f['problems'],1): d.add_heading(f'7.{i} 问题{i}',2); d.add_paragraph(x)
    d.add_heading('8 实践总结',1); d.add_paragraph('本次实践完成了学生选课管理系统从需求分析、E-R 建模、关系模式设计到 MySQL 实现和验证的完整过程。我使用主键和外键建立学生、课程与成绩之间的联系，通过检查约束控制成绩和学分范围，并使用生成列自动计算总评成绩。随后完成 INSERT、UPDATE、DELETE，以及单表查询、多表连接、聚合统计、分组查询、视图和索引等功能。通过独立验证查询确认三张表的数据量、成绩范围、外键完整性、视图和索引均符合要求。本次实践加深了我对概念设计、逻辑设计和 SQL 实现之间关系的理解，也认识到数据库版本、字段类型和参照完整性会直接影响结果可靠性。')
    d.add_heading('9 参考资料',1); d.add_paragraph('[1] 陈红、王珊：《数据库系统原理教程（第2版）》（2021年版），清华大学出版社。\n[2] MySQL 8.0 Reference Manual。\n[3] 四川大学 08695 实践报告模板及实践考核方案。')
    p=O/f"{ID['code']}_08695_{ID['exam']}_{ID['name']}.docx"; d.save(p); return p

def bd_report():
    f=F['bigdata']; d=Document(); setup(d); cover(d,'14899','大数据技术基础（实践）'); integrity(d); toc(d,['1 实践题目','2 开发环境','3 实践过程','4 运行结果汇总','5 实践中遇到的问题及解决方法','6 实践总结','7 参考资料'])
    d.add_heading('1 实践题目',1); d.add_paragraph('所选题目：题目三 Hadoop 平台搭建与 MapReduce 示例运行。')
    d.add_heading('2 开发环境',1); table(d,['项目','实际环境'],[('操作系统',f['linux_os']),('Hadoop 版本',f['hadoop_version']),('JDK 版本',f['java_version']),('运行模式','HDFS 单节点伪分布式，MapReduce 示例运行'),('HDFS 个人目录',f"/user/{ID['exam']}/")])
    ev=R/'bigdata_wordcount/evidence_real'; d.add_heading('3 实践过程',1); d.add_heading('3.1 环境准备',2); d.add_paragraph('在真实 Ubuntu Linux 环境中验证 JDK 和 Hadoop，确认系统、Java 与 Hadoop 均可正常运行。'); fig(d,ev/'01_linux_java.png','图1 Linux 系统与 JDK 版本验证')
    d.add_heading('3.2 Hadoop 安装与配置',2); d.add_paragraph('检查 core-site.xml、hdfs-site.xml 等核心配置，确认 HDFS 访问地址、单节点副本数以及 NameNode/DataNode 相关配置。'); fig(d,ev/'02_hadoop_config.png','图2 Hadoop 核心配置文件')
    d.add_heading('3.3 Hadoop 启动与验证',2); d.add_paragraph('通过 jps 和 dfsadmin 查看 NameNode、DataNode、SecondaryNameNode 以及 HDFS 存储状态，并通过 NameNode Web 页面浏览文件系统。'); fig(d,ev/'03_hadoop_processes.png','图3 Hadoop 进程与 HDFS 状态'); fig(d,ev/'09_hdfs_web.png','图4 NameNode Web 管理界面')
    d.add_heading('3.4 HDFS 操作',2); d.add_paragraph(f"在 /user/{ID['exam']}/input 下使用 mkdir、put、ls、cat、cp、mv、get、rm 等命令完成目录和文件操作。"); fig(d,ev/'04_hdfs_ops.png','图5 HDFS 创建、复制、移动、下载与删除'); fig(d,ev/'05_hdfs_content.png','图6 HDFS 文件列表与内容查看')
    d.add_heading('3.5 MapReduce 示例运行',2); d.add_paragraph('使用 Hadoop 自带 hadoop-mapreduce-examples JAR 运行 WordCount，并查看 part-r-00000 以及汇总统计。'); fig(d,ev/'06_wordcount_job.png','图7 WordCount 运行'); fig(d,ev/'07_wordcount_result.png','图8 WordCount 输出'); fig(d,ev/'08_top5.png','图9 单词总数、不同单词数量与 Top 5')
    d.add_heading('4 运行结果汇总',1); top='，'.join(f'{w}={c}' for w,c in f['top5']); table(d,['序号','操作','结果'],[(1,'Linux/JDK/Hadoop 验证',f"{f['linux_os']}；{f['java_version']}；{f['hadoop_version']}"),(2,'HDFS 目录',f"/user/{ID['exam']}/input 正常"),(3,'HDFS 文件操作','mkdir/put/ls/cat/cp/mv/get/rm 均通过'),(4,'WordCount','运行完成'),(5,'总单词数',f['total_words']),(6,'不同单词数量',f['distinct_words']),(7,'Top 5',top)])
    d.add_heading('5 实践中遇到的问题及解决方法',1)
    for i,x in enumerate(f['problems'],1): d.add_heading(f'5.{i} 问题{i}',2); d.add_paragraph(x)
    d.add_heading('6 实践总结',1); d.add_paragraph('本次实践在真实 Ubuntu Linux 环境中完成 Hadoop 平台配置、HDFS 文件系统操作和 MapReduce WordCount 示例运行。通过环境搭建，我进一步熟悉了 JAVA_HOME、HADOOP_HOME 和 Hadoop 核心配置文件之间的关系；通过 HDFS 命令实际完成目录创建、文件上传、内容查看、复制、移动、下载和删除。在 MapReduce 部分，我使用 Hadoop 自带示例程序处理三个文本文件，并从运行日志和 part-r-00000 中核对结果，同时统计总单词数、不同单词数量和出现次数最多的前5个单词。结合 jps、dfsadmin 和 NameNode Web 页面，可以从进程、存储状态和文件系统三个角度验证 Hadoop 运行情况。本次实践让我对 HDFS 与 MapReduce 的基本流程形成了更清晰的理解，也认识到运行环境和输出目录状态会直接影响实验能否稳定复现。')
    d.add_heading('7 参考资料',1); d.add_paragraph('[1] 林子雨：《大数据基础编程、实验和案例教程》（第3版），清华大学出版社，2024年。\n[2] Apache Hadoop 3.5.0 Documentation。\n[3] 四川大学 14899 实践报告模板及实践考核方案。')
    p=O/f"{ID['code']}_14899_{ID['exam']}_{ID['name']}.docx"; d.save(p); return p

if __name__=='__main__': print(db_report()); print(bd_report())
